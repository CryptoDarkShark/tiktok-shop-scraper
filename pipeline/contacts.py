"""Founder / head of marketing / e-commerce manager for each brand, from public sources only:
the brand's own site (about/team pages) and LinkedIn search results via SerpAPI (1 search per brand).

    python pipeline/contacts.py --batch 2026-09-28 [--apply]   # --apply sets first_name for greetings
    python pipeline/contacts.py hairburst.com

Writes linkedin_contacts.xlsx (all found contacts, with source) and pipeline/contacts_cache.json (cache,
so a brand is never searched twice). Never guesses email addresses.
"""
import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import intl_pipeline as ip  # noqa: E402

CACHE = ip.ROOT / "pipeline" / "contacts_cache.json"
ROLE_RX = re.compile(r"(co-?founder|founder|owner|ceo|managing director|head of marketing|marketing director|"
                     r"marketing manager|brand manager|chief marketing officer|cmo|head of e-?commerce|"
                     r"e-?commerce (?:manager|director|lead)|digital marketing manager|partnerships? manager|"
                     r"influencer (?:marketing )?manager|head of partnerships|head of growth|"
                     r"założyciel(?:ka)?|właściciel(?:ka)?|fondatrice|fondateur|cofondatrice|cofondateur)", re.I)
PRIORITY = ["founder", "co-founder", "owner", "ceo", "managing director", "head of marketing", "marketing director",
            "chief marketing officer", "cmo", "head of e-commerce", "e-commerce manager", "marketing manager",
            "partnerships manager", "influencer", "brand manager", "digital marketing manager", "head of growth"]
SEARCH_NAME = {"Nails.INC": "Nails Inc", "Skin & Tan": "Skinny Tan"}
GL = {"United Kingdom": "uk", "Ireland": "ie", "Poland": "pl", "France": "fr", "Spain": "es", "Italy": "it",
      "Netherlands": "nl", "Belgium": "be"}


def norm(s):
    return re.sub(r"[^a-z0-9]", "", s.lower())


def linkedin_search(b):
    q = (f'site:linkedin.com/in "{SEARCH_NAME.get(b["brand"], b["brand"])}" (founder OR "co-founder" OR "head of marketing" OR '
         f'"marketing manager" OR "ecommerce manager" OR "e-commerce manager")')
    res = ip.serp(q, GL.get(b["country"], "uk"), "en")
    out = []
    for r in res:
        url, title, snip = r.get("link", ""), r.get("title", ""), r.get("snippet", "")
        if "linkedin.com/in/" not in url:
            continue
        # Title format: "Name - Role - Company | LinkedIn" (headline = current role).
        head = title.replace("| LinkedIn", "").strip()
        parts = [x.strip() for x in re.split(r"\s[-–|]\s", head) if x.strip()]
        if not parts:
            continue
        name = parts[0]
        headline = " - ".join(parts[1:])
        # The brand must be in the headline (current role), not only somewhere in the snippet.
        if norm(b["brand"]) not in norm(headline) and norm(b["brand"]) not in norm(title):
            continue
        m = ROLE_RX.search(headline) or ROLE_RX.search(snip if norm(b["brand"]) in norm(snip) else "")
        if norm(b["brand"]) in norm(name):  # a person who happens to be called e.g. "Makear ..."
            continue
        if not m or len(name.split()) > 4 or not re.match(r"^[A-ZÀ-Ž]", name):
            continue
        out.append({"name": name, "role": m.group(0).strip(), "headline": headline[:120], "linkedin": url.split("?")[0],
                    "source": "LinkedIn search result", "evidence": title[:150]})
    return out


def site_people(b):
    """Names stated next to a role on the brand's own pages (about / our story / team)."""
    out = []
    final, html = ip.fetch(b["website"])
    if not html:
        return out
    from bs4 import BeautifulSoup
    home = BeautifulSoup(html, "html.parser")
    for u in ip.find_links(home, final, ["about", "our-story", "story", "team", "founder", "o-nas", "qui-sommes",
                                         "notre-histoire", "chi-siamo", "quienes-somos"], limit=6):
        f, h = ip.fetch(u)
        if not h:
            continue
        t = re.sub(r"\s+", " ", BeautifulSoup(h, "html.parser").get_text(" "))
        for m in re.finditer(r"(?:(?:co-?)?founder|founded by|CEO|head of marketing|założycielk?a?|fondatrice|"
                             r"fondateur)[,:]?\s+(?:and\s+)?([A-Z][a-zà-ž]+(?:\s[A-Z][a-zà-ž'\-]+){1,2})", t):
            name = m.group(1)
            if norm(b["brand"]) in norm(name):
                continue
            out.append({"name": name, "role": ROLE_RX.search(m.group(0)).group(0) if ROLE_RX.search(m.group(0))
                        else "founder", "headline": "", "linkedin": "", "source": f,
                        "evidence": t[max(0, m.start() - 60): m.end() + 40]})
        for m in re.finditer(r"([A-Z][a-zà-ž]+(?:\s[A-Z][a-zà-ž'\-]+){1,2}),?\s+(?:our\s+)?(?:co-?)?founder", t):
            out.append({"name": m.group(1), "role": "founder", "headline": "", "linkedin": "", "source": f,
                        "evidence": t[max(0, m.start() - 40): m.end() + 40]})
    return out


def rank(c):
    r = c["role"].lower()
    return next((i for i, p in enumerate(PRIORITY) if p in r), 99)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("domains", nargs="*")
    ap.add_argument("--batch")
    ap.add_argument("--apply", action="store_true", help="set first_name (greeting) from the best contact")
    a = ap.parse_args()
    brands = ip.curated()
    doms = a.domains or [it["domain"] for it in json.loads((ip.DRAFTS / f"batch-{a.batch}.json").read_text())["items"]]
    cache = ip.load_json(CACHE, {})
    for d in doms:
        b = brands[d]
        if d not in cache:
            found = site_people(b)
            try:
                found += linkedin_search(b)
            except Exception as e:
                print(f"search failed for {b['brand']}: {e} (not cached; re-run to retry)")
                continue
            # dedupe by name; prefer entries with a LinkedIn URL
            best = {}
            for c in found:
                k = norm(c["name"])
                if k not in best or (c["linkedin"] and not best[k]["linkedin"]):
                    best[k] = c
            cache[d] = sorted(best.values(), key=rank)
            ip.DATA.mkdir(parents=True, exist_ok=True)
            CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=1))
    rows, cur = [], json.loads(ip.CURATION.read_text())
    for d in doms:
        b = brands[d]
        people = cache.get(d, [])
        top = people[0] if people and people[0]["name"] and rank(people[0]) < 99 else None
        # "Will and Anna" (two named co-founders) stays whole; otherwise first name only.
        greet = (top["name"] if " and " in top["name"] else top["name"].split()[0]) if top else ""
        for c in people or [{"name": "", "role": "", "linkedin": "", "source": "none found", "evidence": ""}]:
            rows.append([b["brand"], c["name"], c["role"], c["linkedin"], c["source"], c["evidence"],
                         "yes" if greet and c is people[0] else ""])
        if a.apply and d in cur["brands"]:
            cur["brands"][d]["first_name"] = greet
            cur["brands"][d]["first_name_source"] = (people[0]["linkedin"] or people[0]["source"]) if greet else ""
        print(f"{b['brand']}: " + ("; ".join(f"{c['name']} ({c['role']}) {c['linkedin'] or c['source']}"
                                             for c in people) or "none found") + (f"  -> greeting 'Hi {greet},'" if greet else ""))
    if a.apply:
        ip.CURATION.write_text(json.dumps(cur, ensure_ascii=False, indent=1))
    # merge with any existing sheet rows for other brands
    path = ip.ROOT / "linkedin_contacts.xlsx"
    old = []
    if path.exists():
        from openpyxl import load_workbook
        ws = load_workbook(path).active
        old = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)
               if r[0] not in {brands[d]["brand"] for d in doms}]
    ip.write_xlsx(path, ["Brand", "Name", "Role", "LinkedIn URL", "Source", "Evidence", "Used in greeting"],
                  old + rows, title="Contacts")
    print(f"-> {path}")


if __name__ == "__main__":
    main()
